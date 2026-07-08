#!/usr/bin/env python3
"""Compare code changes against Excel field specification."""
import io
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
EXCEL = Path(r"e:\Chrome\汇率及关联字段改动 (1).xlsx")

# Excel J-column doc -> code file mapping (product confirmed hardcode scope)
J_DOC_CODE_MAP = {
    ("WMS系统", "销售出库单"): [
        "SoOutstockServiceImpl.java",
        "SoOutstockDetailServiceImpl.java",
        "SyncB2CSoOutstockServiceImpl.java",
    ],
    ("WMS系统", "退货通知单"): [
        "SoReturnNoticeServiceImpl.java",
        "SoReturnNoticeDetailServiceImpl.java",
    ],
    ("WMS系统", "退货签收单"): [
        "SoReturnReceiveServiceImpl.java",
        "SoReturnReceiveDetailServiceImpl.java",
    ],
    ("WMS系统", "退货入库单"): [
        "SoReturnInstockServiceImpl.java",
        "SoReturnInstockDetailServiceImpl.java",
        "RestCloudPlatformNewReturnInstockConsumerService.java",
        "SyncSoReturnServiceImpl.java",
    ],
    ("OMS系统", "B2B销售订单"): [
        "SoDetailServiceImpl.java",
        "SoInfoServiceImpl.java",
    ],
    ("OMS系统", "B2C-销售订单"): [],  # only exchange rate in J; no dedicated hardcode file in whitelist
    ("模具管理", "展会订单"): ["ExhibitionOrderDetailServiceImpl.java"],
    ("物流管理", "SKU成本"): ["InventorySkuCostServiceImpl.java"],
    ("物流管理", "头程对账单"): [
        "TmsFirstMileReconciliationServiceImpl.java",
        "TmsFirstMileReconciliationDetailServiceImpl.java",
    ],
    ("物流管理", "物流暂估账单"): ["FirstMileEstimatedBillServiceImpl.java"],
    ("物流管理", "尾程费用"): ["LogisticsBillCostServiceImpl.java"],
    ("物流管理", "小包费用分摊"): ["SmallBagCostAllocationServiceImpl.java"],
    ("物流管理", "中转费用分摊"): ["TransferDeclareServiceImpl.java"],
    ("物流管理", "B2C报关对账单"): [
        "TmsB2cDeclareReconciliationServiceImpl.java",
        "TmsB2cDeclareReconciliationDetailServiceImpl.java",
    ],
    ("物流管理", "自发货费用"): ["LogisticsLargeServiceImpl.java"],
    ("物流管理", "运费模板"): ["ShippingCalculationServiceImpl.java"],
}

# Excel field Chinese -> Java field
FIELD_MAP = {
    "汇率": "exchange_rate",
    "直接汇率": "direct_exchange_rate / exchange_rate",
    "销售单价(本位币)": "price_lc / priceLc",
    "销售单价(本位币）": "price_lc / priceLc",
    "含税单价(本位币)": "tax_price_lc / taxPriceLc",
    "价税合计(本位币)": "all_amount_local_currency / allAmountLocalCurrency",
    "退货金额（本位币）": "return_amount_local_currency / returnAmountLocalCurrency",
    "含税退货金额（本位币）": "tax_return_amount_local_currency / taxReturnAmountLocalCurrency",
}

# K-column storage-only fields we additionally changed (not in J whitelist script)
K_EXTRA_CODE = {
    ("OMS系统", "B2B售后"): [
        "SoReturnServiceImpl.java",
        "SoReturnDetailServiceImpl.java",
        "NewPlatformB2bReturnOrderConsumerService.java",
    ],
}

PATCH39_FILES = set()
rules_path = ROOT / "scripts/patch_excel39_hardcode.py"
text = rules_path.read_text(encoding="utf-8")
for line in text.splitlines():
    if line.strip().startswith('"erp-server/'):
        PATCH39_FILES.add(line.split('"')[1].split("/")[-1])


def load_excel():
    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for row in ws.iter_rows(values_only=True):
        if any(cell is not None and str(cell).strip() for cell in row):
            rows.append([("" if c is None else str(c).strip()) for c in row])
    items = []
    cols = [
        "system", "domain", "doc", "owner", "field", "source",
        "fe_before", "be_before", "fe_change", "be_hardcode", "be_storage",
        "history", "formula",
    ]
    for r in rows[1:]:
        while len(r) < 13:
            r.append("")
        items.append(dict(zip(cols, r[:13])))
    return items


def git_changed_server_files():
    out = subprocess.check_output(
        ["git", "diff", "--name-only", "--", "erp-server"],
        cwd=ROOT,
        text=True,
    )
    return {Path(p).name for p in out.splitlines() if p.strip()}


def main():
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    items = load_excel()
    hardcode = [i for i in items if i["be_hardcode"]]
    changed = git_changed_server_files()

    print("=" * 70)
    print("Excel J列硬编码 39 条 vs 代码改动核对")
    print("=" * 70)

    covered = []
    missing_code = []
    for i, x in enumerate(hardcode, 1):
        key = (x["system"], x["doc"])
        code_files = J_DOC_CODE_MAP.get(key, [])
        java_field = FIELD_MAP.get(x["field"], x["field"])
        hit = [f for f in code_files if f in changed]
        status = "OK" if hit or (x["field"] == "汇率" and not code_files) else (
            "PARTIAL" if code_files else "NO_MAP"
        )
        if hit:
            covered.append(x)
        elif code_files:
            missing_code.append((x, code_files))
        print(f"{i:02d}. [{status:7}] {x['system']} / {x['doc']} / {x['field']}")
        print(f"       Java: {java_field}")
        print(f"       J列: {x['be_hardcode']} | K列: {x['be_storage']}")
        if code_files:
            print(f"       期望文件: {', '.join(code_files)}")
            print(f"       已改动: {', '.join(hit) if hit else '无'}")
        print()

    print("=" * 70)
    print("J列仅「汇率」、patch39白名单未覆盖的单据（预期靠DB+全局，无硬编码改动）")
    print("=" * 70)
    rate_only = [x for x in hardcode if x["field"] in ("汇率", "直接汇率")
                 and J_DOC_CODE_MAP.get((x["system"], x["doc"]), None) == []]
    for x in rate_only:
        print(f"- {x['system']} / {x['doc']} / {x['field']}")

    print()
    print("=" * 70)
    print("K列存储六位 - 退货本位币 & OMS B2B售后（J列无硬编码或仅K列）")
    print("=" * 70)
    ret_fields = [i for i in items if "本位币" in i["field"] and "退货" in i["field"]]
    for x in ret_fields:
        key = (x["system"], x["doc"])
        extra = K_EXTRA_CODE.get(("OMS系统", "B2B售后"), [])
        wms_files = J_DOC_CODE_MAP.get(key, [])
        all_expected = list(dict.fromkeys(wms_files + extra))
        hit = [f for f in all_expected if f in changed]
        print(f"- {x['doc']} / {x['field']} | J:{x['be_hardcode'] or '-'} | 改动:{', '.join(hit) or '透传/上游'}")

    oms_b2b = [i for i in items if i["doc"] == "B2B售后" and "本位币" in i["field"]]
    if oms_b2b:
        print("\nOMS B2B售后:")
        for x in oms_b2b:
            hit = [f for f in K_EXTRA_CODE[("OMS系统", "B2B售后")] if f in changed]
            print(f"  {x['field']} | J:{x['be_hardcode'] or '-'} | K:{x['be_storage']} | 改动:{', '.join(hit)}")

    print()
    print("=" * 70)
    print("erp-server 已改动但不在 Excel J列39 / K列补改范围")
    print("=" * 70)
    expected_all = set(PATCH39_FILES)
    for files in J_DOC_CODE_MAP.values():
        expected_all.update(files)
    for files in K_EXTRA_CODE.values():
        expected_all.update(files)
    # extra return-related
    expected_all.update([
        "SoReturnReceiveServiceImpl.java",
        "SoReturnReceiveDetailServiceImpl.java",
        "RestCloudPlatformNewReturnInstockConsumerService.java",
        "SyncSoReturnServiceImpl.java",
        "SoReturnServiceImpl.java",
        "SoReturnDetailServiceImpl.java",
        "NewPlatformB2bReturnOrderConsumerService.java",
    ])
    extra_changed = sorted(changed - expected_all)
    for f in extra_changed:
        print(f"- {f}")

    print()
    print("=" * 70)
    print("patch_excel39 白名单文件改动情况")
    print("=" * 70)
    for f in sorted(PATCH39_FILES):
        mark = "YES" if f in changed else "NO"
        print(f"  [{mark}] {f}")

    print()
    print("SUMMARY")
    print(f"  Excel J列硬编码: {len(hardcode)}")
    print(f"  J列有代码映射且已改动单据字段: {len(covered)} / {len(hardcode)}")
    print(f"  erp-server额外改动文件数: {len(extra_changed)}")


if __name__ == "__main__":
    main()
